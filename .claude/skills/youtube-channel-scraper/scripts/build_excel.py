"""Build <ChannelName>_Video_Report.xlsx with two sheets:
  - "All Videos": every scraped video, newest first.
  - "Matched Criteria": only is_match videos, highest relevance_score first (capped at top_n).
"""
import re
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=11)
BODY_FONT = Font(name="Arial", size=10)
LINK_FONT = Font(name="Arial", size=10, color="0563C1", underline="single")
WRAP = Alignment(wrap_text=True, vertical="top")
CELL_LIMIT = 5000  # Sheets/Excel cap cells at 32-50k chars; long descriptions would break the upload

ALL_VIDEOS_COLUMNS = [
    ("Title", 45), ("Video Link", 40), ("Published Date", 14), ("Type", 8), ("Views", 10),
    ("Likes", 10), ("Comments", 10), ("Duration", 10), ("Matched?", 10),
    ("Relevance Score", 14), ("Matched Keywords", 22), ("Description", 70),
]
MATCHED_COLUMNS = [
    ("Relevance Score", 14), ("Title", 45), ("Video Link", 40), ("Published Date", 14), ("Type", 8),
    ("Views", 10), ("Likes", 10), ("Comments", 10), ("Duration", 10),
    ("Matched Keywords", 22), ("Transcript Hits", 60), ("Description", 70),
]


def iso8601_duration_to_hms(dur):
    if not dur:
        return ""
    m = re.match(r"P(?:(\d+)D)?T?(?:(\d+)H)?(?:(\d+)M)?(?:(\d+)S)?$", dur)
    if not m:
        return dur
    d, h, mi, s = (int(x) if x else 0 for x in m.groups())
    h += d * 24
    return f"{h}:{mi:02d}:{s:02d}" if h else f"{mi}:{s:02d}"


def style_header(ws, columns):
    for idx, (name, width) in enumerate(columns, 1):
        cell = ws.cell(row=1, column=idx, value=name)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 22


def fmt_date(published_at):
    try:
        return datetime.strptime((published_at or "")[:10], "%Y-%m-%d").strftime("%Y-%m-%d")
    except ValueError:
        return published_at or ""


def clip(text):
    text = text or ""
    return text if len(text) <= CELL_LIMIT else text[:CELL_LIMIT] + "..."


def write_row(ws, row_idx, values, link_col_idx):
    for col_idx, val in enumerate(values, 1):
        cell = ws.cell(row=row_idx, column=col_idx, value=val)
        cell.font = BODY_FONT
        cell.alignment = WRAP
        if col_idx == link_col_idx and val:
            cell.hyperlink = val
            cell.font = LINK_FONT
    ws.row_dimensions[row_idx].height = 15


def build_all_videos_sheet(ws, videos):
    style_header(ws, ALL_VIDEOS_COLUMNS)
    for i, v in enumerate(videos, start=2):
        write_row(ws, i, [
            v.get("title", ""), v.get("url", ""), fmt_date(v.get("published_at")), v.get("video_type", ""),
            v.get("view_count", 0), v.get("like_count"), v.get("comment_count"),
            iso8601_duration_to_hms(v.get("duration")),
            "Yes" if v.get("is_match") else "No", v.get("relevance_score"),
            ", ".join(v.get("matched_keywords", [])), clip(v.get("description")),
        ], link_col_idx=2)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(ALL_VIDEOS_COLUMNS))}{len(videos) + 1}"


def build_matched_sheet(ws, videos):
    style_header(ws, MATCHED_COLUMNS)
    for i, v in enumerate(videos, start=2):
        write_row(ws, i, [
            v.get("relevance_score"), v.get("title", ""), v.get("url", ""), fmt_date(v.get("published_at")),
            v.get("video_type", ""), v.get("view_count", 0), v.get("like_count"), v.get("comment_count"),
            iso8601_duration_to_hms(v.get("duration")),
            ", ".join(v.get("matched_keywords", [])),
            clip("\n".join(v.get("transcript_snippets", []))), clip(v.get("description")),
        ], link_col_idx=3)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(MATCHED_COLUMNS))}{len(videos) + 1}"


def build_excel(channel_display_name, videos, out_dir, top_n=30):
    all_sorted = sorted(videos, key=lambda v: v.get("published_at") or "", reverse=True)
    matched = sorted([v for v in videos if v.get("is_match")],
                     key=lambda v: v.get("relevance_score") or 0, reverse=True)
    total_matched = len(matched)
    if top_n:
        matched = matched[:top_n]

    wb = Workbook()
    ws1 = wb.active
    ws1.title = "All Videos"
    build_all_videos_sheet(ws1, all_sorted)
    build_matched_sheet(wb.create_sheet("Matched Criteria"), matched)

    safe_name = re.sub(r"[^A-Za-z0-9_-]+", "_", channel_display_name).strip("_") or "channel"
    out_path = Path(out_dir) / f"{safe_name}_Video_Report.xlsx"
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    print(f"Wrote {len(all_sorted)} rows to 'All Videos' and top {len(matched)} of "
          f"{total_matched} matches to 'Matched Criteria' in {out_path}")
    return out_path
