"""Scan a recalculated OTB workbook for formula errors and print the numbers Review mode needs.

Usage: python check_otb.py <workbook.xlsx>
Run recalc_excel.ps1 first — without cached values every formula reads as None.
"""
import json
import sys

from openpyxl import load_workbook

ERRORS = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!")


def table(ws, header_row, first_col="A"):
    """Read rows under a header until the first blank first-column cell."""
    headers = [c.value for c in ws[header_row] if c.value is not None]
    out, r = [], header_row + 1
    while ws[f"{first_col}{r}"].value not in (None, ""):
        out.append(dict(zip(headers, [ws.cell(row=r, column=j + 1).value for j in range(len(headers))])))
        r += 1
    return out


def find_row(ws, text, col=1):
    for row in ws.iter_rows(min_col=col, max_col=col):
        if row[0].value == text:
            return row[0].row
    raise ValueError(f"'{text}' not found on {ws.title}")


def main(path):
    wb = load_workbook(path, data_only=True)
    errors, empty_formulas = [], 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value in ERRORS:
                    errors.append(f"{ws.title}!{c.coordinate} {c.value}")

    formulas = load_workbook(path)
    for ws in formulas.worksheets:
        if ws.title in ("Sales Log", "Stock Aging"):  # blank log rows legitimately evaluate to ""
            continue
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("=") and wb[ws.title][c.coordinate].value is None:
                    empty_formulas += 1

    db = wb["Dashboard"]
    cat_head = find_row(db, "Category")
    src_head = find_row(db, "Source")
    age_head = find_row(db, "Bucket")
    cash_start = find_row(db, "Working capital")
    cash = {}
    r = cash_start
    while db[f"A{r}"].value:
        cash[db[f"A{r}"].value] = db[f"B{r}"].value
        r += 1

    inputs = wb["Inputs"]
    mix_check = inputs.cell(row=find_row(inputs, "Total"), column=3).value

    report = {
        "file": path,
        "formula_errors": errors[:50],
        "total_errors": len(errors),
        "formulas_without_cached_value": empty_formulas,
        "as_of": str(db["A2"].value),
        "current_month": str(db["B4"].value),
        "plan_month_index": db["B5"].value,
        "mix_check": mix_check,
        "categories": [row for row in table(db, cat_head) if row.get("Category") != "Total"],
        "sourcing": table(db, src_head),
        "cash": cash,
        "aging": table(db, age_head),
    }
    print(json.dumps(report, indent=2, default=str))


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit("Usage: python check_otb.py <workbook.xlsx>")
    main(sys.argv[1])
